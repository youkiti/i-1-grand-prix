import sys
from pathlib import Path

# srcディレクトリをパスに追加
sys.path.append(str(Path(__file__).parent.parent / "src"))

from interview_analysis.loader import read_file_content, load_documents_from_folder

def test_file_loading():
    test_dir = Path("test_data")
    test_dir.mkdir(exist_ok=True)

    # 1. Create dummy files
    (test_dir / "test.txt").write_text("This is a test text file.", encoding="utf-8")
    (test_dir / "test.csv").write_text("col1,col2\nval1,val2", encoding="utf-8")
    
    # PDF, Excel, Word files are harder to create purely with text, 
    # so we will just test the text-based ones and the folder loading logic for now.
    # If pypdf/openpyxl/python-docx are installed, we could try to create them using those libs if we wanted to be thorough.
    
    try:
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Header1", "Header2"])
        ws.append(["Data1", "Data2"])
        wb.save(test_dir / "test.xlsx")
        print("Created test.xlsx")
    except ImportError:
        print("openpyxl not installed, skipping xlsx creation")

    try:
        from docx import Document
        doc = Document()
        doc.add_paragraph("This is a test docx file.")
        doc.save(test_dir / "test.docx")
        print("Created test.docx")
    except ImportError:
        print("python-docx not installed, skipping docx creation")

    # 2. Test read_file_content
    print("\n--- Testing read_file_content ---")
    for file_path in test_dir.iterdir():
        print(f"Reading {file_path.name}...")
        content = read_file_content(file_path)
        print(f"Content length: {len(content)}")
        print(f"Preview: {content[:50]}...")

    # 3. Test load_documents_from_folder
    print("\n--- Testing load_documents_from_folder ---")
    combined = load_documents_from_folder(test_dir)
    print(f"Combined content length: {len(combined)}")
    print("Combined content preview:")
    print(combined[:200])

    # 4. Test iter_documents_from_folder
    print("\n--- Testing iter_documents_from_folder ---")
    from interview_analysis.loader import iter_documents_from_folder
    count = 0
    for filename, content in iter_documents_from_folder(test_dir):
        count += 1
        print(f"Iterated: {filename} (len={len(content)})")
    print(f"Total iterated files: {count}")

    # Cleanup
    # import shutil
    # shutil.rmtree(test_dir)

if __name__ == "__main__":
    test_file_loading()
